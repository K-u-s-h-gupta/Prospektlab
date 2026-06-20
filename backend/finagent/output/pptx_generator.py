"""PPTX report generator for deep research deliverables."""

from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Any

from finagent.output.contracts import AnswerArtifact

logger = logging.getLogger("finagent.output.pptx_generator")


class PPTXGenerator:
    """Generate McKinsey-style slide decks from an AnswerArtifact."""

    def __init__(self) -> None:
        pass

    def generate(self, artifact: AnswerArtifact, output_path: Path | str) -> Path:
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        except ImportError as exc:
            raise ImportError("python-pptx is required for PPTX export. Install with: pip install python-pptx") from exc

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # Title slide
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12), Inches(1.5))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = artifact.title or "Research Report"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 58, 95)

        subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12), Inches(0.8))
        stf = subtitle_box.text_frame
        sp = stf.paragraphs[0]
        sp.text = f"ProspektLab Deep Research | {artifact.analyst_view.thesis if artifact.analyst_view else ''}"
        sp.font.size = Pt(16)
        sp.font.color.rgb = RGBColor(100, 116, 139)

        # Executive summary
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_heading(slide, "Executive Summary")
        summary_text = artifact.summary or "No summary available."
        self._add_bullet_box(slide, summary_text, Inches(0.5), Inches(1.2), Inches(12), Inches(5.5))

        # Metric cards
        if artifact.metric_cards:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._add_heading(slide, "Financial Snapshot")
            x, y = 0.5, 1.2
            for i, mc in enumerate(artifact.metric_cards[:6]):
                col = i % 3
                row = i // 3
                self._add_metric_card(
                    slide,
                    label=mc.label,
                    value=mc.value,
                    change=mc.change,
                    x=Inches(0.5 + col * 4.0),
                    y=Inches(1.2 + row * 2.0),
                    width=Inches(3.5),
                    height=Inches(1.5),
                )

        # Sections
        for section in artifact.sections:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._add_heading(slide, section.title)
            content = section.body or ""
            self._add_bullet_box(slide, content, Inches(0.5), Inches(1.2), Inches(12), Inches(5.5))

        # Citations
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        self._add_heading(slide, "Sources & Citations")
        citation_text = "\n".join(
            f"• {c.kind}: {c.url}" for c in artifact.citations[:10]
        ) if artifact.citations else "No citations available."
        self._add_bullet_box(slide, citation_text, Inches(0.5), Inches(1.2), Inches(12), Inches(5.5))

        out = Path(output_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(out))
        logger.info("[PPTX] generated %s", out)
        return out

    def generate_bytes(self, artifact: AnswerArtifact) -> bytes:
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
            path = self.generate(artifact, tmp.name)
            data = path.read_bytes()
            path.unlink()
            return data

    def _add_heading(self, slide: Any, text: str) -> None:
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
        tf = box.text_frame
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = RGBColor(30, 58, 95)

    def _add_bullet_box(self, slide: Any, text: str, x: float, y: float, w: float, h: float) -> None:
        from pptx.util import Pt
        box = slide.shapes.add_textbox(x, y, w, h)
        tf = box.text_frame
        tf.word_wrap = True
        for i, line in enumerate(text.split("\n")):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(14)
            p.space_after = Pt(8)

    def _add_metric_card(self, slide: Any, label: str, value: str, change: str | None, x: float, y: float, width: float, height: float) -> None:
        from pptx.util import Pt, Inches
        from pptx.dml.color import RGBColor
        shape = slide.shapes.add_shape(1, x, y, width, height)  # 1 = rectangle
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(247, 250, 252)
        shape.line.color.rgb = RGBColor(226, 232, 240)

        tf = shape.text_frame
        tf.margin_top = Inches(0.1)
        tf.margin_left = Inches(0.1)
        p1 = tf.paragraphs[0]
        p1.text = label.upper()
        p1.font.size = Pt(10)
        p1.font.color.rgb = RGBColor(113, 128, 150)

        p2 = tf.add_paragraph()
        p2.text = str(value)
        p2.font.size = Pt(20)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(26, 32, 44)

        if change:
            p3 = tf.add_paragraph()
            p3.text = str(change)
            p3.font.size = Pt(12)
            p3.font.color.rgb = RGBColor(16, 185, 129) if "+" in str(change) or "up" in str(change).lower() else RGBColor(239, 68, 68)


def generate_pptx(artifact: AnswerArtifact, output_path: Path | str) -> Path:
    return PPTXGenerator().generate(artifact, output_path)


__all__ = ["PPTXGenerator", "generate_pptx"]
