"""Server-side report exporter for FinAgent.

Generates professional analyst-style documents (PDF, DOCX, XLSX)
from an AnswerArtifact. PDF uses headless Chromium via patchright.
Document export is available for deep-mode sessions only.

Layout is modeled on institutional research report formats (HSIE,
Motilal Oswal, Kotak, Jefferies): cover block with rating/TP card,
executive summary, analyst view, financial snapshot, sectional
deep-dive, embedded charts, valuation, peer comparison, citations,
disclaimer footer.
"""
from __future__ import annotations

import base64
import io
import logging
import re
from datetime import date
from typing import Literal

from finagent.output.compliance import CHAT_DISCLAIMER
from finagent.output.contracts import AnswerArtifact, ResponseDepth

logger = logging.getLogger("finagent.report_exporter")

Style = Literal["sell_side", "buy_side", "auto"]


def can_generate_documents(depth: ResponseDepth | str) -> bool:
    """Document export is only available for deep-mode research."""
    if isinstance(depth, str):
        return depth.lower() == "deep"
    return depth == ResponseDepth.DEEP


# ── Helpers ────────────────────────────────────────────────────────────

def _esc(text: object) -> str:
    s = str(text or "")
    return (s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


_INVALID_SHEET_CHARS = re.compile(r"[\\/?*\[\]:]")


def _safe_sheet_name(raw: str, used: set[str], fallback: str = "Sheet") -> str:
    """openpyxl rejects: \\ / ? * [ ] : and names >31 chars or duplicates."""
    cleaned = _INVALID_SHEET_CHARS.sub(" ", str(raw or "")).strip()[:31] or fallback
    base = cleaned
    i = 2
    while cleaned.lower() in used:
        suffix = f" ({i})"
        cleaned = (base[: 31 - len(suffix)] + suffix)
        i += 1
    used.add(cleaned.lower())
    return cleaned


def _format_stance(stance: str) -> tuple[str, str]:
    """Map stance to (label, color) like HSIE BUY/ADD/REDUCE/SELL."""
    s = (stance or "neutral").lower()
    mapping = {
        "positive": ("BUY", "#10B981"),
        "watch": ("ADD", "#3B82F6"),
        "neutral": ("HOLD", "#6B7280"),
        "negative": ("REDUCE", "#EF4444"),
    }
    return mapping.get(s, ("HOLD", "#6B7280"))


# ── HTML template (used for PDF) ───────────────────────────────────────

_REPORT_CSS = """
@page { size: A4; margin: 18mm 16mm 22mm 16mm; }
* { box-sizing: border-box; }
body { font-family: 'Inter', 'Segoe UI', Arial, sans-serif; font-size: 10.5px; line-height: 1.55; color: #1a202c; margin: 0; }
.cover { display: grid; grid-template-columns: 1fr 200px; gap: 16px; border-bottom: 3px solid #1e3a5f; padding-bottom: 14px; margin-bottom: 16px; }
.cover-main h1 { font-size: 22px; font-weight: 800; color: #1e3a5f; margin: 4px 0 6px 0; line-height: 1.2; }
.cover-main .thesis { font-size: 13px; color: #2c5282; font-style: italic; margin: 0 0 10px 0; font-weight: 500; }
.cover-main .brand-row { display: flex; align-items: center; gap: 8px; font-size: 10px; color: #4a5568; }
.cover-main .brand-logo { width: 18px; height: 18px; background: #1e3a5f; color: #fff; border-radius: 4px; display: inline-flex; align-items: center; justify-content: center; font-weight: 700; font-size: 11px; }
.cover-main .brand-name { font-weight: 700; color: #1e3a5f; letter-spacing: 0.06em; text-transform: uppercase; font-size: 10px; }
.cover-main .meta { font-size: 9.5px; color: #718096; letter-spacing: 0.04em; margin-top: 6px; }
.cover-main .summary { font-size: 11px; color: #2d3748; margin-top: 8px; }
.rating-card { background: #f7fafc; border: 1px solid #cbd5e0; border-radius: 6px; padding: 10px; font-size: 9.5px; }
.rating-card .stance { font-size: 16px; font-weight: 800; margin-bottom: 4px; }
.rating-card .rc-row { display: flex; justify-content: space-between; padding: 3px 0; border-bottom: 1px dotted #e2e8f0; }
.rating-card .rc-row:last-child { border-bottom: none; }
.rating-card .rc-row .lbl { color: #718096; }
.rating-card .rc-row .val { font-weight: 600; color: #1a202c; }
h2 { font-size: 13px; color: #1e3a5f; background: #edf2f7; padding: 5px 8px; margin: 16px 0 8px 0; border-left: 4px solid #2c5282; }
h3 { font-size: 11.5px; color: #2d3748; margin: 12px 0 4px 0; font-weight: 700; }
p { margin: 4px 0 8px 0; }
ul, ol { margin: 4px 0 8px 18px; padding: 0; }
li { margin-bottom: 3px; }
table { border-collapse: collapse; width: 100%; margin: 6px 0 12px 0; font-size: 9.5px; }
th { background: #2c5282; color: #fff; padding: 5px 7px; text-align: left; font-weight: 600; font-size: 9px; }
td { padding: 4px 7px; border-bottom: 1px solid #e2e8f0; }
tr:nth-child(even) td { background: #f9fafb; }
td.num { text-align: right; font-variant-numeric: tabular-nums; }
.snapshot-table th, .snapshot-table td { font-size: 9px; }
.metric-grid { display: grid; grid-template-columns: repeat(4, 1fr); gap: 6px; margin: 8px 0 12px 0; }
.metric-card { background: #f7fafc; border: 1px solid #e2e8f0; border-radius: 5px; padding: 7px 9px; }
.metric-label { font-size: 8.5px; color: #718096; text-transform: uppercase; letter-spacing: 0.04em; }
.metric-value { font-size: 13px; font-weight: 700; color: #1a202c; margin-top: 1px; }
.metric-change { font-size: 9px; margin-top: 1px; }
.positive { color: #10B981; }
.negative { color: #EF4444; }
.chart-block { margin: 8px 0 14px 0; page-break-inside: avoid; }
.chart-img { max-width: 100%; height: auto; border: 1px solid #e2e8f0; border-radius: 4px; }
.analyst-view { background: linear-gradient(90deg,#edf2f7 0%,#f7fafc 100%); border-left: 4px solid #1e3a5f; padding: 9px 12px; margin: 10px 0 14px 0; border-radius: 0 5px 5px 0; }
.analyst-view h3 { margin-top: 0; color: #1e3a5f; }
.av-row { display: flex; gap: 14px; font-size: 10px; margin-bottom: 4px; }
.av-row .lbl { color: #718096; font-weight: 600; text-transform: uppercase; letter-spacing: 0.04em; font-size: 9px; }
.citation { font-size: 8.5px; color: #4a5568; line-height: 1.45; margin: 2px 0; }
.citation a { color: #2c5282; text-decoration: none; }
.rating-legend { background: #f7fafc; border: 1px solid #e2e8f0; border-radius: 5px; padding: 8px 10px; margin-top: 12px; font-size: 9px; color: #4a5568; }
.disclaimer { font-size: 8px; color: #718096; margin-top: 18px; border-top: 1px solid #e2e8f0; padding-top: 6px; line-height: 1.45; }
.footer-brand { font-size: 8px; color: #a0aec0; text-align: center; margin-top: 4px; letter-spacing: 0.08em; text-transform: uppercase; }
"""


def _artifact_to_report_html(artifact: AnswerArtifact, style: Style = "sell_side") -> str:
    """Convert an AnswerArtifact to analyst-report HTML suitable for print."""
    av = artifact.analyst_view
    stance_label, stance_color = _format_stance(av.stance if av else "neutral")
    today_str = date.today().strftime("%d %B %Y")

    parts: list[str] = [
        "<!DOCTYPE html><html><head><meta charset='utf-8'>",
        f"<style>{_REPORT_CSS}</style></head><body>",
    ]

    # ── Cover block ────────────────────────────────────────────────────
    parts.append('<div class="cover">')
    parts.append('<div class="cover-main">')
    parts.append('<div class="brand-row"><span class="brand-logo">P</span><span class="brand-name">ProspektLab Research</span></div>')
    parts.append(f'<h1>{_esc(artifact.title or "Research Report")}</h1>')
    if av and av.thesis:
        parts.append(f'<p class="thesis">{_esc(av.thesis)}</p>')
    parts.append(f'<div class="meta">{today_str}  |  Deep Research Report</div>')
    if artifact.summary:
        parts.append(f'<p class="summary">{_esc(artifact.summary)}</p>')
    parts.append("</div>")

    # Right sidebar: rating card
    parts.append('<div class="rating-card">')
    parts.append(f'<div class="stance" style="color:{stance_color}">{stance_label}</div>')
    if av:
        parts.append(f'<div class="rc-row"><span class="lbl">Conviction</span><span class="val">{_esc(av.conviction.upper())}</span></div>')
        if av.horizon:
            parts.append(f'<div class="rc-row"><span class="lbl">Horizon</span><span class="val">{_esc(av.horizon)}</span></div>')
        if av.valuation_summary:
            vs = av.valuation_summary
            if vs.range_low and vs.range_high:
                parts.append(f'<div class="rc-row"><span class="lbl">Target Range</span><span class="val">{_esc(vs.range_low)}–{_esc(vs.range_high)}</span></div>')
            if vs.method:
                parts.append(f'<div class="rc-row"><span class="lbl">Method</span><span class="val">{_esc(vs.method)}</span></div>')
    parts.append(f'<div class="rc-row"><span class="lbl">Date</span><span class="val">{today_str}</span></div>')
    parts.append("</div>")
    parts.append("</div>")  # /cover

    # ── Key metric grid ────────────────────────────────────────────────
    if artifact.metric_cards:
        parts.append('<div class="metric-grid">')
        for mc in artifact.metric_cards[:8]:
            change_cls = "positive" if mc.change_type == "positive" else ("negative" if mc.change_type == "negative" else "")
            parts.append('<div class="metric-card">')
            parts.append(f'<div class="metric-label">{_esc(mc.label)}</div>')
            parts.append(f'<div class="metric-value">{_esc(mc.value)}</div>')
            if mc.change:
                parts.append(f'<div class="metric-change {change_cls}">{_esc(mc.change)}</div>')
            parts.append("</div>")
        parts.append("</div>")

    # ── Analyst view box ───────────────────────────────────────────────
    if av:
        parts.append('<div class="analyst-view">')
        parts.append("<h3>Analyst View</h3>")
        parts.append(f'<div class="av-row"><span><span class="lbl">Stance</span> {stance_label}</span> <span><span class="lbl">Conviction</span> {_esc(av.conviction.upper())}</span>')
        if av.horizon:
            parts.append(f' <span><span class="lbl">Horizon</span> {_esc(av.horizon)}</span>')
        parts.append("</div>")
        if av.thesis:
            parts.append(f"<p><strong>Thesis:</strong> {_esc(av.thesis)}</p>")
        if av.key_risks:
            parts.append("<p><strong>Key Risks:</strong></p><ul>")
            for r in av.key_risks[:6]:
                parts.append(f"<li>{_esc(r)}</li>")
            parts.append("</ul>")
        if av.valuation_summary and (av.valuation_summary.range_low or av.valuation_summary.assumptions):
            vs = av.valuation_summary
            parts.append(f'<p><strong>Valuation ({_esc(vs.method or "blend")}):</strong> {_esc(vs.range_low)} – {_esc(vs.range_high)}</p>')
            if vs.assumptions:
                parts.append(f'<p style="font-size:9.5px;color:#4a5568"><em>{_esc(vs.assumptions)}</em></p>')
        parts.append("</div>")

    # ── Sections (narrative) ───────────────────────────────────────────
    for section in artifact.sections:
        parts.append(f"<h2>{_esc(section.title)}</h2>")
        for blk in _parse_md_blocks(section.body or ""):
            kind = blk[0]
            if kind == "p":
                parts.append(f"<p>{_inline_html(blk[1])}</p>")
            elif kind == "h":
                parts.append(f"<h3>{_inline_html(blk[2])}</h3>")
            elif kind == "ul":
                parts.append("<ul>")
                for it in blk[1]:
                    parts.append(f"<li>{_inline_html(it)}</li>")
                parts.append("</ul>")
            elif kind == "ol":
                parts.append("<ol>")
                for it in blk[1]:
                    parts.append(f"<li>{_inline_html(it)}</li>")
                parts.append("</ol>")
            elif kind == "table":
                header, rows = blk[1], blk[2]
                parts.append('<table><thead><tr>')
                for col in header:
                    parts.append(f"<th>{_inline_html(col)}</th>")
                parts.append("</tr></thead><tbody>")
                for row in rows:
                    parts.append("<tr>")
                    for j, cell in enumerate(row):
                        cls = "num" if j > 0 and _looks_numeric(cell) else ""
                        parts.append(f'<td class="{cls}">{_inline_html(cell)}</td>')
                    parts.append("</tr>")
                parts.append("</tbody></table>")

    # ── Charts (embedded PNGs) ─────────────────────────────────────────
    for chart in artifact.charts:
        if chart.png_data:
            parts.append('<div class="chart-block">')
            parts.append(f"<h3>{_esc(chart.title)}</h3>")
            if chart.subtitle:
                parts.append(f'<p style="font-size:9.5px;color:#718096;margin-top:0">{_esc(chart.subtitle)}</p>')
            parts.append(f'<img class="chart-img" src="data:image/png;base64,{chart.png_data}" alt="{_esc(chart.title)}">')
            parts.append("</div>")

    # ── Tables (financial / comparison) ────────────────────────────────
    for table in artifact.tables:
        parts.append(f"<h3>{_esc(table.title)}</h3>")
        parts.append('<table class="snapshot-table"><thead><tr>')
        for col in table.columns:
            parts.append(f"<th>{_esc(col)}</th>")
        parts.append("</tr></thead><tbody>")
        for row in table.rows:
            parts.append("<tr>")
            for j, cell in enumerate(row):
                cls = "num" if j > 0 and _looks_numeric(cell) else ""
                parts.append(f'<td class="{cls}">{_esc(cell)}</td>')
            parts.append("</tr>")
        parts.append("</tbody></table>")

    # ── Citations ──────────────────────────────────────────────────────
    if artifact.citations:
        parts.append("<h2>Sources & References</h2>")
        for c in artifact.citations[:20]:
            domain = c.domain or c.url
            parts.append(f'<p class="citation">• <strong>{_esc(c.title)}</strong> — <a href="{_esc(c.url)}">{_esc(domain)}</a></p>')

    # ── Rating legend ──────────────────────────────────────────────────
    parts.append('<div class="rating-legend"><strong>Stance Legend:</strong>')
    parts.append(' BUY (positive thesis, conviction with timeframe) · ADD (constructive watch) · HOLD (balanced) · REDUCE (negative thesis).')
    parts.append(" Stance is an analytical reading of the evidence; not a personalised investment recommendation.</div>")

    # ── Disclaimer ─────────────────────────────────────────────────────
    parts.append(f'<div class="disclaimer">{_esc(CHAT_DISCLAIMER)}</div>')
    parts.append('<div class="footer-brand">ProspektLab Research · prospektlab.com</div>')
    parts.append("</body></html>")
    return "".join(parts)


_MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_MD_LINK_RE = re.compile(r"\[([^\]]+)\]\((https?://[^)]+)\)")
_MD_HEADING_RE = re.compile(r"^(#{2,6})\s+(.+)$")
_MD_BULLET_RE = re.compile(r"^[-*•]\s+(.+)$")
_MD_NUM_RE = re.compile(r"^\d+[.)]\s+(.+)$")


def _inline_segments(text: str) -> list[tuple[str, bool]]:
    """Split prose into (chunk, is_bold) honoring **bold**; markdown links flattened to label."""
    text = _MD_LINK_RE.sub(r"\1", text or "")
    segs: list[tuple[str, bool]] = []
    pos = 0
    for m in _MD_BOLD_RE.finditer(text):
        if m.start() > pos:
            segs.append((text[pos:m.start()], False))
        segs.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        segs.append((text[pos:], False))
    return segs or [(text, False)]


def _inline_html(text: str) -> str:
    """Escape text and re-apply **bold** as <strong>; flatten markdown links to label."""
    text = _MD_LINK_RE.sub(r"\1", text or "")
    out: list[str] = []
    pos = 0
    for m in _MD_BOLD_RE.finditer(text):
        if m.start() > pos:
            out.append(_esc(text[pos:m.start()]))
        out.append(f"<strong>{_esc(m.group(1))}</strong>")
        pos = m.end()
    if pos < len(text):
        out.append(_esc(text[pos:]))
    return "".join(out)


def _is_table_line(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.count("|") >= 2


def _is_separator_row(cells: list[str]) -> bool:
    return bool(cells) and all(set(c) <= set("-: ") and "-" in c for c in cells)


def _split_table_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def _parse_md_blocks(body: str) -> list[tuple]:
    """Parse a section body into typed blocks so markdown renders properly in DOCX/PDF.

    Block tuples: ("p", text) | ("h", level, text) | ("ul", [items]) |
    ("ol", [items]) | ("table", header, rows).
    """
    lines = (body or "").replace("\r\n", "\n").split("\n")
    blocks: list[tuple] = []
    para: list[str] = []
    n = len(lines)
    i = 0

    def flush() -> None:
        nonlocal para
        if para:
            text = " ".join(s.strip() for s in para if s.strip()).strip()
            if text:
                blocks.append(("p", text))
            para = []

    while i < n:
        line = lines[i].strip()
        if not line:
            flush()
            i += 1
            continue
        m = _MD_HEADING_RE.match(line)
        if m:
            flush()
            blocks.append(("h", len(m.group(1)), m.group(2).strip()))
            i += 1
            continue
        if _is_table_line(line):
            rows = []
            while i < n and _is_table_line(lines[i]):
                rows.append(_split_table_row(lines[i]))
                i += 1
            if rows:
                flush()
                header = rows[0]
                data = [r for r in rows[1:] if not _is_separator_row(r)]
                blocks.append(("table", header, data))
            continue
        if _MD_BULLET_RE.match(line):
            items = []
            while i < n and _MD_BULLET_RE.match(lines[i].strip()):
                items.append(_MD_BULLET_RE.match(lines[i].strip()).group(1).strip())
                i += 1
            flush()
            blocks.append(("ul", items))
            continue
        if _MD_NUM_RE.match(line):
            items = []
            while i < n and _MD_NUM_RE.match(lines[i].strip()):
                items.append(_MD_NUM_RE.match(lines[i].strip()).group(1).strip())
                i += 1
            flush()
            blocks.append(("ol", items))
            continue
        para.append(line)
        i += 1
    flush()
    return blocks


def _looks_numeric(text: object) -> bool:
    s = str(text or "").strip().replace(",", "").replace("₹", "").replace("$", "").replace("%", "").replace("(", "-").replace(")", "")
    if not s or s.lower() in ("n/a", "na", "-", "—"):
        return False
    try:
        float(s)
        return True
    except ValueError:
        return False


# ── PDF export via headless Chromium ───────────────────────────────────

async def export_pdf(artifact: AnswerArtifact, style: Style = "sell_side") -> bytes:
    """Generate a professional PDF from the artifact using headless Chromium."""
    html = _artifact_to_report_html(artifact, style)

    try:
        from patchright.async_api import async_playwright
    except ImportError as exc:
        raise RuntimeError(
            "patchright is required for PDF export. "
            "Install with: pip install patchright && patchright install chromium"
        ) from exc

    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        try:
            page = await browser.new_page()
            # `load` (not `networkidle`) is correct for inline-only HTML —
            # `networkidle` can hang for ~30s waiting for the 500ms idle window
            # when there are no network requests to settle.
            await page.set_content(html, wait_until="load", timeout=20000)
            pdf_bytes = await page.pdf(
                format="A4",
                print_background=True,
                margin={"top": "16mm", "bottom": "20mm", "left": "14mm", "right": "14mm"},
                prefer_css_page_size=True,
            )
        finally:
            await browser.close()
    return pdf_bytes


# ── DOCX export — HSIE-style structured layout ─────────────────────────

def export_docx(artifact: AnswerArtifact, style: Style = "sell_side") -> bytes:
    """Generate a DOCX analyst report from the artifact, mirroring institutional
    research report structure: cover block, exec summary, analyst view box,
    metric snapshot, sections, charts, tables, citations, disclaimer."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError as exc:
        raise RuntimeError("python-docx is required for DOCX export. pip install python-docx") from exc

    av = artifact.analyst_view
    stance_label, stance_hex = _format_stance(av.stance if av else "neutral")
    today_str = date.today().strftime("%d %B %Y")

    doc = Document()

    # Set narrow margins for institutional research look
    for section in doc.sections:
        section.top_margin = Cm(1.6)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    def _set_cell_bg(cell, hex_color: str) -> None:
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def _add_h1(text: str, color_hex: str = "1E3A5F") -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor.from_string(color_hex)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    def _add_h2(text: str) -> None:
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string("2D3748")
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)

    def _add_md_paragraph(text: str, style: str | None = None):
        p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
        for chunk, bold in _inline_segments(text):
            run = p.add_run(chunk)
            if bold:
                run.bold = True
        return p

    def _emit_docx_table(columns: list, rows: list) -> None:
        if not columns:
            return
        tbl = doc.add_table(rows=1, cols=len(columns))
        tbl.style = "Light Grid Accent 1"
        tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
        hdr_cells = tbl.rows[0].cells
        for i, col_name in enumerate(columns):
            cell = hdr_cells[i]
            _set_cell_bg(cell, "2C5282")
            cp = cell.paragraphs[0]
            cr = cp.add_run(str(col_name))
            cr.bold = True
            cr.font.size = Pt(9)
            cr.font.color.rgb = RGBColor.from_string("FFFFFF")
        for row_data in rows:
            row_cells = tbl.add_row().cells
            for i, cell_val in enumerate(row_data):
                if i >= len(row_cells):
                    break
                c = row_cells[i]
                c.text = ""
                cp = c.paragraphs[0]
                cr = cp.add_run(str(cell_val))
                cr.font.size = Pt(9)
                if i > 0 and _looks_numeric(cell_val):
                    cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # ── Cover header table (title left, rating card right) ─────────────
    cover_tbl = doc.add_table(rows=1, cols=2)
    cover_tbl.autofit = False
    cover_tbl.columns[0].width = Cm(13)
    cover_tbl.columns[1].width = Cm(5)
    left_cell = cover_tbl.rows[0].cells[0]
    right_cell = cover_tbl.rows[0].cells[1]

    # Left: brand + title + thesis + meta + summary
    brand_p = left_cell.paragraphs[0]
    brand_run = brand_p.add_run("ProspektLab Research")
    brand_run.bold = True
    brand_run.font.size = Pt(9)
    brand_run.font.color.rgb = RGBColor.from_string("1E3A5F")

    title_p = left_cell.add_paragraph()
    title_run = title_p.add_run(artifact.title or "Research Report")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor.from_string("1E3A5F")

    if av and av.thesis:
        thesis_p = left_cell.add_paragraph()
        thesis_run = thesis_p.add_run(av.thesis)
        thesis_run.italic = True
        thesis_run.font.size = Pt(11)
        thesis_run.font.color.rgb = RGBColor.from_string("2C5282")

    meta_p = left_cell.add_paragraph()
    meta_run = meta_p.add_run(f"{today_str}  ·  Deep Research Report")
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor.from_string("718096")

    if artifact.summary:
        sum_p = left_cell.add_paragraph()
        sum_run = sum_p.add_run(artifact.summary)
        sum_run.font.size = Pt(10)
        sum_p.paragraph_format.space_before = Pt(4)

    # Right: rating card
    _set_cell_bg(right_cell, "F7FAFC")
    rc_p = right_cell.paragraphs[0]
    rc_run = rc_p.add_run(stance_label)
    rc_run.bold = True
    rc_run.font.size = Pt(18)
    rc_run.font.color.rgb = RGBColor.from_string(stance_hex.replace("#", ""))

    def _rc_row(label: str, val: str) -> None:
        p = right_cell.add_paragraph()
        r1 = p.add_run(label)
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor.from_string("718096")
        r2 = p.add_run(f"  {val}")
        r2.bold = True
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor.from_string("1A202C")

    if av:
        _rc_row("Conviction", av.conviction.upper())
        if av.horizon:
            _rc_row("Horizon", av.horizon)
        if av.valuation_summary and av.valuation_summary.range_low:
            vs = av.valuation_summary
            _rc_row("Target", f"{vs.range_low}–{vs.range_high}")
            if vs.method:
                _rc_row("Method", vs.method)
    _rc_row("Date", today_str)

    # Divider
    doc.add_paragraph("─" * 90).runs[0].font.color.rgb = RGBColor.from_string("CBD5E0")

    # ── Key metrics snapshot (compact table) ───────────────────────────
    if artifact.metric_cards:
        _add_h1("Key Metrics Snapshot")
        cards = artifact.metric_cards[:8]
        cols = min(4, len(cards))
        rows_needed = (len(cards) + cols - 1) // cols
        mtbl = doc.add_table(rows=rows_needed * 2, cols=cols)
        mtbl.style = "Light Grid Accent 1"
        for idx, mc in enumerate(cards):
            r = (idx // cols) * 2
            c = idx % cols
            label_cell = mtbl.rows[r].cells[c]
            value_cell = mtbl.rows[r + 1].cells[c]
            lbl_p = label_cell.paragraphs[0]
            lbl_run = lbl_p.add_run(mc.label.upper())
            lbl_run.font.size = Pt(7.5)
            lbl_run.font.color.rgb = RGBColor.from_string("718096")
            val_p = value_cell.paragraphs[0]
            val_run = val_p.add_run(str(mc.value))
            val_run.bold = True
            val_run.font.size = Pt(11)
            if mc.change:
                ch_run = val_p.add_run(f"  {mc.change}")
                ch_run.font.size = Pt(8)
                ch_color = "10B981" if mc.change_type == "positive" else ("EF4444" if mc.change_type == "negative" else "718096")
                ch_run.font.color.rgb = RGBColor.from_string(ch_color)

    # ── Analyst view box ───────────────────────────────────────────────
    if av:
        _add_h1("Analyst View")
        av_tbl = doc.add_table(rows=1, cols=1)
        av_cell = av_tbl.rows[0].cells[0]
        _set_cell_bg(av_cell, "EDF2F7")
        if av.thesis:
            t_p = av_cell.paragraphs[0]
            t_p.add_run("Thesis: ").bold = True
            t_p.add_run(av.thesis)
        if av.key_risks:
            k_p = av_cell.add_paragraph()
            k_p.add_run("Key Risks:").bold = True
            for r in av.key_risks[:6]:
                bp = av_cell.add_paragraph(r, style="List Bullet")
                bp.paragraph_format.left_indent = Cm(0.6)
        if av.valuation_summary and av.valuation_summary.assumptions:
            vp = av_cell.add_paragraph()
            vp.add_run("Assumptions: ").bold = True
            vp.add_run(av.valuation_summary.assumptions)

    # ── Narrative sections ─────────────────────────────────────────────
    for section in artifact.sections:
        _add_h1(section.title)
        for blk in _parse_md_blocks(section.body or ""):
            kind = blk[0]
            if kind == "p":
                _add_md_paragraph(blk[1])
            elif kind == "h":
                _add_h2(blk[2])
            elif kind == "ul":
                for it in blk[1]:
                    _add_md_paragraph(it, style="List Bullet")
            elif kind == "ol":
                for it in blk[1]:
                    _add_md_paragraph(it, style="List Number")
            elif kind == "table":
                _emit_docx_table(blk[1], blk[2])

    # ── Embedded charts ────────────────────────────────────────────────
    for chart in artifact.charts:
        if chart.png_data:
            _add_h2(chart.title)
            if chart.subtitle:
                sub_p = doc.add_paragraph()
                sub_run = sub_p.add_run(chart.subtitle)
                sub_run.italic = True
                sub_run.font.size = Pt(9)
                sub_run.font.color.rgb = RGBColor.from_string("718096")
            try:
                img_bytes = base64.b64decode(chart.png_data)
                img_stream = io.BytesIO(img_bytes)
                doc.add_picture(img_stream, width=Inches(6.0))
            except Exception:
                logger.warning("Failed to embed chart %s in DOCX", chart.id, exc_info=True)

    # ── Financial / comparison tables ──────────────────────────────────
    for table_data in artifact.tables:
        _add_h2(table_data.title)
        _emit_docx_table(table_data.columns, table_data.rows)

    # ── Citations ──────────────────────────────────────────────────────
    if artifact.citations:
        _add_h1("Sources & References")
        for c in artifact.citations[:20]:
            p = doc.add_paragraph(style="List Bullet")
            r1 = p.add_run(f"{c.title} ")
            r1.bold = True
            r1.font.size = Pt(9)
            r2 = p.add_run(f"— {c.domain or c.url}")
            r2.font.size = Pt(8.5)
            r2.font.color.rgb = RGBColor.from_string("4A5568")

    # ── Rating legend + disclaimer ─────────────────────────────────────
    leg_p = doc.add_paragraph()
    leg_p.paragraph_format.space_before = Pt(12)
    leg_run = leg_p.add_run(
        "Stance Legend: BUY (positive thesis with timeframe) · ADD (constructive watch) · "
        "HOLD (balanced) · REDUCE (negative thesis). "
        "Stance is an analytical reading of the evidence; not a personalised investment recommendation."
    )
    leg_run.font.size = Pt(8)
    leg_run.font.color.rgb = RGBColor.from_string("4A5568")

    disc_p = doc.add_paragraph()
    disc_run = disc_p.add_run(CHAT_DISCLAIMER)
    disc_run.font.size = Pt(8)
    disc_run.font.color.rgb = RGBColor.from_string("718096")

    brand_p = doc.add_paragraph()
    brand_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    br_run = brand_p.add_run("PROSPEKTLAB RESEARCH  ·  prospektlab.com")
    br_run.font.size = Pt(7.5)
    br_run.font.color.rgb = RGBColor.from_string("A0AEC0")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


# ── XLSX export — multi-sheet workbook with sanitized names ────────────

def export_xlsx(artifact: AnswerArtifact) -> bytes:
    """Generate an Excel workbook with metrics, tables, and chart series.
    Sheet names are sanitized to avoid openpyxl InvalidWorksheetName errors."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for XLSX export. pip install openpyxl") from exc

    wb = Workbook()
    used: set[str] = set()

    header_fill = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    title_font = Font(bold=True, size=14, color="1E3A5F")
    subtitle_font = Font(italic=True, size=10, color="2C5282")
    label_font = Font(bold=True, size=9, color="718096")
    thin = Side(border_style="thin", color="CBD5E0")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)

    # ── Summary sheet ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = _safe_sheet_name("Summary", used)
    ws["A1"] = "ProspektLab Research"
    ws["A1"].font = Font(bold=True, size=10, color="1E3A5F")
    ws["A2"] = artifact.title or "Research Report"
    ws["A2"].font = title_font
    if artifact.analyst_view and artifact.analyst_view.thesis:
        ws["A3"] = artifact.analyst_view.thesis
        ws["A3"].font = subtitle_font
    ws["A4"] = date.today().strftime("%d %B %Y")
    ws["A4"].font = Font(size=9, color="718096")

    row = 6
    ws.cell(row=row, column=1, value="EXECUTIVE SUMMARY").font = label_font
    row += 1
    ws.cell(row=row, column=1, value=artifact.summary or "").alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 90
    ws.column_dimensions["A"].width = 90

    # Analyst view block
    if artifact.analyst_view:
        av = artifact.analyst_view
        stance_label, _ = _format_stance(av.stance)
        row += 2
        ws.cell(row=row, column=1, value="ANALYST VIEW").font = label_font
        row += 1
        ws.cell(row=row, column=1, value=f"Stance: {stance_label}   ·   Conviction: {av.conviction.upper()}").font = Font(bold=True, size=10)
        if av.thesis:
            row += 1
            ws.cell(row=row, column=1, value=f"Thesis: {av.thesis}").alignment = Alignment(wrap_text=True)
        if av.key_risks:
            row += 1
            ws.cell(row=row, column=1, value="Key Risks:").font = Font(bold=True)
            for r in av.key_risks[:6]:
                row += 1
                ws.cell(row=row, column=1, value=f"  • {r}").alignment = Alignment(wrap_text=True)

    # ── Metrics sheet ──────────────────────────────────────────────────
    if artifact.metric_cards:
        ws_m = wb.create_sheet(_safe_sheet_name("Metrics", used))
        ws_m.append(["Metric", "Value", "Change", "Type"])
        for cell in ws_m[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        for mc in artifact.metric_cards:
            ws_m.append([mc.label, mc.value, mc.change, mc.change_type])
        ws_m.column_dimensions["A"].width = 32
        ws_m.column_dimensions["B"].width = 22
        ws_m.column_dimensions["C"].width = 14
        ws_m.column_dimensions["D"].width = 12

    # ── Data tables ────────────────────────────────────────────────────
    for i, table_data in enumerate(artifact.tables):
        sheet_name = _safe_sheet_name(table_data.title or f"Table {i+1}", used, fallback=f"Table{i+1}")
        ws_t = wb.create_sheet(sheet_name)
        if table_data.columns:
            ws_t.append(list(table_data.columns))
            for cell in ws_t[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
            for row_data in table_data.rows:
                ws_t.append([str(x) for x in row_data])
            for col_idx in range(1, len(table_data.columns) + 1):
                col_letter = ws_t.cell(row=1, column=col_idx).column_letter
                max_len = max(
                    [len(str(table_data.columns[col_idx - 1]))]
                    + [len(str(row[col_idx - 1])) for row in table_data.rows if col_idx - 1 < len(row)],
                    default=10,
                )
                ws_t.column_dimensions[col_letter].width = min(max(max_len + 2, 10), 38)

    # ── Chart series data ──────────────────────────────────────────────
    for j, chart in enumerate(artifact.charts):
        if chart.series:
            sheet_name = _safe_sheet_name(chart.title or f"Chart Data {j+1}", used, fallback=f"Chart{j+1}")
            ws_c = wb.create_sheet(sheet_name)
            header_row = ["Label"] + [s.name for s in chart.series]
            ws_c.append(header_row)
            for cell in ws_c[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.border = border
            for idx, label in enumerate(chart.labels):
                row = [label]
                for s in chart.series:
                    val = s.data[idx].value if idx < len(s.data) else ""
                    row.append(val)
                ws_c.append(row)
            ws_c.column_dimensions["A"].width = 18

    # ── Citations sheet ────────────────────────────────────────────────
    if artifact.citations:
        ws_s = wb.create_sheet(_safe_sheet_name("Sources", used))
        ws_s.append(["Title", "Domain", "URL", "Kind"])
        for cell in ws_s[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
        for c in artifact.citations[:50]:
            ws_s.append([c.title, c.domain or "", c.url, c.kind])
        ws_s.column_dimensions["A"].width = 50
        ws_s.column_dimensions["B"].width = 24
        ws_s.column_dimensions["C"].width = 60

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
